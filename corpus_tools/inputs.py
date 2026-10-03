import csv
import hashlib
import io
from pathlib import Path
import re
import shutil
import uuid

from .storage import atomic_json, now

from .storage import identity

MEDIA = {".mp4", ".mkv", ".mov", ".avi", ".flv", ".wmv", ".mpg", ".webm",
         ".mp3", ".wav", ".flac", ".m4a", ".ogg"}
SUPPORTED = MEDIA | {".srt", ".txt", ".csv", ".xlsx"}
TIME = re.compile(r"(\d{1,3}):(\d{2}):(\d{2})[,.](\d{3})")


def decode(raw):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("无法识别文本编码，请另存为 UTF-8。")


def millis(value):
    match = TIME.fullmatch(str(value).strip())
    if not match:
        raise ValueError(f"无效字幕时间：{value}")
    h, m, s, ms = map(int, match.groups())
    if m >= 60 or s >= 60:
        raise ValueError("字幕分钟、秒数超出范围。")
    return ((h * 60 + m) * 60 + s) * 1000 + ms


def timestamp(value):
    h, rem = divmod(int(value), 3600000)
    m, rem = divmod(rem, 60000)
    s, ms = divmod(rem, 1000)
    return f"{h:02}:{m:02}:{s:02},{ms:03}"


def parse_srt(text):
    result = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()):
        lines = block.strip().splitlines()
        if not lines:
            continue
        timing = 1 if lines[0].strip().isdigit() else 0
        if len(lines) <= timing + 1 or "-->" not in lines[timing]:
            raise ValueError(f"字幕块格式错误：{lines[0][:40]}")
        left, right = lines[timing].split("-->", 1)
        start, end = millis(left.strip()), millis(right.strip().split()[0])
        if end < start:
            raise ValueError("字幕结束时间早于开始时间。")
        sentence = " ".join(line.strip() for line in lines[timing+1:]).strip()
        if sentence:
            result.append({"text": sentence, "start_ms": start, "end_ms": end,
                           "cue": len(result) + 1})
    if not result:
        raise ValueError("没有可读取的字幕。")
    return result


def tabular_rows(file, raw):
    if file.suffix.lower() == ".csv":
        return list(csv.DictReader(io.StringIO(decode(raw))))
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("读取 XLSX 需要安装 openpyxl。") from None
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    rows = []
    try:
        for sheet in book.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else "" for v in next(iterator, ())]
            for values in iterator:
                row = {key: "" if value is None else str(value) for key, value in zip(headers, values)}
                row["_sheet"] = sheet.title
                rows.append(row)
    finally:
        book.close()
    return rows


def pick(row, names, default=""):
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default


def media_fingerprint(file):
    info = file.stat()
    return identity(info.st_size, info.st_mtime_ns)



def text_records(file, raw):
    if file.suffix.lower() == ".srt":
        return parse_srt(decode(raw))
    if file.suffix.lower() == ".txt":
        result = [{"text": line.strip(), "cue": i + 1, "start_ms": None, "end_ms": None}
                  for i, line in enumerate(decode(raw).splitlines()) if line.strip()]
    elif file.suffix.lower() in (".csv", ".xlsx"):
        result = []
        for i, row in enumerate(tabular_rows(file, raw)):
            text = pick(row, ("文本", "内容", "text", "transcript", "sentence", "台词", "语料"))
            if not text:
                if any(str(v or "").strip() for k, v in row.items() if k != "_sheet"):
                    raise ValueError("表格需要 文本/内容/text/transcript/sentence/台词/语料 列。")
                continue
            start = pick(row, ("开始时间", "start", "start_time"))
            end = pick(row, ("结束时间", "end", "end_time"))
            start_ms, end_ms = millis(start) if start else None, millis(end) if end else None
            if start_ms is not None and end_ms is not None and end_ms < start_ms:
                raise ValueError("表格中有结束时间早于开始时间的记录。")
            result.append({"text": text, "cue": i+1, "start_ms": start_ms, "end_ms": end_ms,
                           "imported_fields": row})
    else:
        raise ValueError("已有文本支持 SRT、TXT、CSV、XLSX。")
    if not result:
        raise ValueError("输入文件没有有效文本。")
    return result


def atomic_text(file, value):
    temporary = file.with_name("." + file.name + "-" + uuid.uuid4().hex + ".part")
    try:
        temporary.write_text(value, encoding="utf-8")
        temporary.replace(file)
    finally:
        temporary.unlink(missing_ok=True)


def save_transcript(item, records, origin, extra=None):
    """媒体、SRT、TXT 同层保存；替换文本时留存旧副本。调用者持有 task.lock。"""
    if not records:
        raise ValueError("没有有效转写文本。")
    media = item.read()["media"]
    stem = Path(media["filename"]).stem
    txt = item.path(stem + ".txt")
    srt = item.path(stem + ".srt")
    timed = all(row.get("start_ms") is not None and row.get("end_ms") is not None for row in records)
    history = item.path(".corpus/history/" + uuid.uuid4().hex)
    previous = {}
    for file in (txt, srt):
        if file.exists():
            history.mkdir(parents=True, exist_ok=True)
            backup = history / file.name
            shutil.copy2(file, backup)
            previous[file] = backup
    try:
        atomic_text(txt, "\n".join(row["text"] for row in records) + "\n")
        if timed:
            atomic_text(srt, compose_srt(records))
        else:
            srt.unlink(missing_ok=True)
        with item.edit() as data:
            sid = media["id"]
            data["segments"] = {identity(sid, row["cue"]): {
                "id": identity(sid, row["cue"]), "source_id": sid, **row} for row in records}
            data["sources"][sid].update(transcribed=True, transcript=srt.name if timed else txt.name,
                                       text_file=txt.name, text_origin=origin, error="", **(extra or {}))
            data["hits"] = []
            data["search"]["performed"] = False
    except Exception:
        for file in (txt, srt):
            if file in previous:
                shutil.copy2(previous[file], file)
            else:
                file.unlink(missing_ok=True)
        raise


def attach_text(item, file):
    """已有文本替换该媒体的检索文本，不把 SRT 与 TXT 重复计入。"""
    file = Path(file).resolve()
    raw = file.read_bytes()
    records = text_records(file, raw)
    with item.job():
        snapshot = item.path(".corpus/inputs/" + hashlib.sha256(raw).hexdigest()[:24] + file.suffix.lower())
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        snapshot.write_bytes(raw)
        save_transcript(item, records, "import", {"text_source": file.name})
    return len(records)


def compose_srt(records):
    return "\n".join(f"{i}\n{timestamp(row['start_ms'])} --> {timestamp(row['end_ms'])}\n{row['text']}\n"
                     for i, row in enumerate(records, 1))
