import csv
from datetime import datetime
from pathlib import Path
import uuid

from .inputs import timestamp
from .storage import atomic_json, now

HEADERS = ["来源文件", "序号", "开始时间", "结束时间", "关键词", "匹配文本",
           "本段第几次命中", "左侧上下文", "右侧上下文", "文本", "原始文本",
           "起始字符位置", "结束字符位置"]


def table_values(row):
    return [row["source_name"], row["cue"],
            timestamp(row["start_ms"]) if row["start_ms"] is not None else "",
            timestamp(row["end_ms"]) if row["end_ms"] is not None else "",
            row["keyword"], row["matched_term"], row["occurrence"],
            row["left_context"], row["right_context"], row["text"], row["original_text"],
            row["offset"], row["end_offset"]]


def excel_literal(value):
    # CSV 由 Excel 打开时避免把语料执行为公式；原文仍保留在 JSON/XLSX。
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export(item, xlsx=True):
    data = item.read()
    if not data["search"].get("performed"):
        raise ValueError("请先检索关键词；新增输入或转写后需重新检索。")
    rows = item.rows(data=data)
    workbook_type = None
    if xlsx:
        try:
            from openpyxl import Workbook
            workbook_type = Workbook
        except ImportError:
            raise RuntimeError("导出 XLSX 需要安装 openpyxl；也可使用仅 CSV 模式。") from None
    tag = datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8]
    folder = item.root
    prefix = Path(data["media"]["filename"]).stem + "-检索-" + tag
    csv_file = item.path(prefix + ".csv")
    xlsx_file = item.path(prefix + ".xlsx")
    json_file = item.path(".corpus/exports/" + prefix + ".json")
    manifest_file = item.path(".corpus/exports/" + prefix + ".manifest.json")
    manifest = {"state": "writing", "created_at": now(), "media": data["media"]["filename"],
                "revision": data["revision"], "search": data["search"], "rows": len(rows)}
    atomic_json(manifest_file, manifest)
    with csv_file.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows([excel_literal(v) for v in table_values(row)] for row in rows)
    atomic_json(json_file, rows)
    if workbook_type:
        from openpyxl.styles import Alignment, Font, PatternFill
        book = workbook_type()
        sheet = book.active
        sheet.title = "检索结果"
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(table_values(row))
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="156F75")
        widths = [28, 10, 18, 18, 16, 16, 18, 40, 40, 60, 60, 18, 18]
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        book.save(xlsx_file)
        book.close()
    manifest["state"] = "complete"
    manifest["files"] = [csv_file.name] + ([xlsx_file.name] if xlsx else [])
    atomic_json(manifest_file, manifest)
    return {"folder": str(folder), "rows": len(rows), "csv": str(csv_file),
            "xlsx": str(xlsx_file) if xlsx else None, "json": str(json_file),
            "manifest": str(manifest_file)}
